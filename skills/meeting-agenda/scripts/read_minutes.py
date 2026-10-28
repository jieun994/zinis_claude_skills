#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""meeting-minutes 표준 양식 회의록(.xlsx)을 읽어 JSON으로 출력한다.

Claude는 엑셀 바이너리를 직접 읽을 수 없으므로, 직전 회의록에서 이월 항목을
가져올 때 이 스크립트를 거친다.

사용법:
    python read_minutes.py <회의록.xlsx> [-o out.json] [--sheet 시트명]

출력 JSON:
{
  "source": "회의록.xlsx",
  "sheet": "2026.07.20",
  "overview": {"project": "...", "회의명": "...", "일시": "...",
               "작성자": "...", "참석자": [...], "근거문서": [...]},
  "discussions": [{"category","items","content","note"}],
  "todos": [{"category","side","content","request_date","due_date","done_date","assignee"}],
  "issues": [{"category","side","content","status","occur_date","plan_date","done_date"}],
  "carryover": {
    "open_todos":  [...],   # 완료일이 비어 있는 TO-DO
    "open_issues": [...],   # 현황이 완료/종료가 아닌 이슈
    "undecided":   [...]    # 미정 / 확인 필요 / TBD 표기가 남은 항목
  }
}

carryover 세 묶음이 다음 회의 안건의 1순위 후보다.
표준 양식이 아니면 섹션을 못 찾고 에러로 알린다 — 컬럼을 추측하지 않는다.
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

COLS = "BCDEFGH"

SEC_DISCUSSION = "회의 내용"
SEC_TODO = "TO-DO"
SEC_ISSUE = "ISSUE"

# 이슈가 닫혔다고 볼 수 있는 현황 표현
CLOSED_STATUS = ("완료", "종료", "해결", "처리완료", "closed", "done")

# 값이 확정되지 않았거나 다음 회의에서 되짚어야 함을 뜻하는 표현
UNDECIDED_PAT = re.compile(
    r"미정|확인\s*필요|TBD|보류|재논의|추후\s*결정|"
    r"검토\s*중|검토\s*후|검토하겠|내부\s*검토|내부\s*논의|이해\s*차이",
    re.IGNORECASE,
)


def err(msg):
    print(msg, file=sys.stderr)
    sys.exit(1)


def cell(ws, col, row):
    v = ws[f"{col}{row}"].value
    if v is None:
        return ""
    return str(v).strip()


def row_empty(ws, row):
    return all(not cell(ws, c, row) for c in COLS)


def find_sections(ws):
    """B열을 훑어 섹션 시작 행을 찾는다. {섹션명: 헤더행}"""
    found = {}
    for row in range(1, ws.max_row + 1):
        b = cell(ws, "B", row)
        if not b:
            continue
        if SEC_DISCUSSION in b and SEC_DISCUSSION not in found:
            found[SEC_DISCUSSION] = row
        elif SEC_TODO in b.upper() and SEC_TODO not in found:
            found[SEC_TODO] = row
        elif SEC_ISSUE in b.upper() and SEC_ISSUE not in found:
            found[SEC_ISSUE] = row
    return found


def read_data_rows(ws, start, stop):
    """start부터 stop 직전까지, 내용이 있는 행만 [(행번호, {열:값})]로 반환."""
    out = []
    for row in range(start, stop):
        if row_empty(ws, row):
            continue
        out.append((row, {c: cell(ws, c, row) for c in COLS}))
    return out


def read_overview(ws, end_row):
    """회의 개요 영역(B2 ~ 회의 내용 직전)을 라벨-값으로 읽는다.

    B열에 라벨이 있고 C열에 값이 있는 구조. 라벨 없이 C열만 있는 행은
    바로 위 라벨의 이어지는 값(참석자 2행, 근거문서 4행)으로 붙인다.
    """
    ov = {"project": "", "회의명": "", "일시": "", "작성자": "",
          "참석자": [], "근거문서": []}

    title = cell(ws, "B", 2)
    m = re.match(r"\[(.+?)\]", title)
    if m:
        ov["project"] = m.group(1).strip()

    label = None
    for row in range(3, end_row):
        b = cell(ws, "B", row)
        c = cell(ws, "C", row)
        if b:
            label = re.sub(r"\s*\(.*?\)\s*", "", b.replace("\n", " ")).strip()
        if not c or not label:
            continue
        parts = [p.strip() for p in c.split("\n") if p.strip()]
        if "참석" in label:
            ov["참석자"].extend(parts)
        elif "근거" in label or "문서" in label:
            ov["근거문서"].extend(parts)
        elif "회의명" in label:
            ov["회의명"] = ov["회의명"] or c
        elif "일시" in label or "일자" in label:
            ov["일시"] = ov["일시"] or c
        elif "작성" in label:
            ov["작성자"] = ov["작성자"] or c
    return ov


def is_undecided(*values):
    return any(UNDECIDED_PAT.search(v) for v in values if v)


def main():
    ap = argparse.ArgumentParser(description="회의록 엑셀 → JSON")
    ap.add_argument("xlsx", help="meeting-minutes 표준 양식 회의록 파일")
    ap.add_argument("-o", "--output", default=None, help="JSON 저장 경로 (없으면 stdout)")
    ap.add_argument("--sheet", default=None, help="시트명 (없으면 첫 시트)")
    args = ap.parse_args()

    path = Path(args.xlsx)
    if not path.exists():
        err(f"파일을 찾을 수 없음: {path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        err(f"엑셀(.xlsx) 파일이 아님: {path.suffix} — 이 형식은 직접 읽어서 처리하세요.")

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        err(f"엑셀을 열 수 없음 (손상/암호 가능): {e}")

    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]

    sec = find_sections(ws)
    if SEC_DISCUSSION not in sec:
        err("표준 회의록 양식이 아님 — B열에서 '회의 내용' 섹션을 찾지 못했습니다. "
            "컬럼 의미를 추측하지 않고 멈춥니다. 파일 구조를 확인하세요.")

    disc_hdr = sec[SEC_DISCUSSION] + 1          # 헤더행(구분/항목/내용/비고)
    todo_hdr = sec.get(SEC_TODO)
    issue_hdr = sec.get(SEC_ISSUE)

    disc_end = todo_hdr or issue_hdr or (ws.max_row + 1)
    disc_rows = read_data_rows(ws, disc_hdr + 1, disc_end)

    todo_rows = []
    if todo_hdr:
        todo_end = issue_hdr or (ws.max_row + 1)
        todo_rows = read_data_rows(ws, todo_hdr + 2, todo_end)

    issue_rows = []
    if issue_hdr:
        issue_rows = read_data_rows(ws, issue_hdr + 2, ws.max_row + 1)

    discussions = [
        {"row": r, "category": v["B"], "items": v["C"],
         "content": v["D"], "note": v["E"]}
        for r, v in disc_rows
    ]
    todos = [
        {"row": r, "category": v["B"], "side": v["C"], "content": v["D"],
         "request_date": v["E"], "due_date": v["F"],
         "done_date": v["G"], "assignee": v["H"]}
        for r, v in todo_rows
    ]
    issues = [
        {"row": r, "category": v["B"], "side": v["C"], "content": v["D"],
         "status": v["E"], "occur_date": v["F"],
         "plan_date": v["G"], "done_date": v["H"]}
        for r, v in issue_rows
    ]

    open_todos = [
        t for t in todos
        if not t["done_date"] or UNDECIDED_PAT.search(t["done_date"])
    ]
    open_issues = [
        s for s in issues
        if not any(k in s["status"].lower() for k in CLOSED_STATUS)
    ]
    undecided = []
    for t in todos:
        if is_undecided(t["assignee"], t["due_date"], t["content"]):
            undecided.append({"type": "todo", **t})
    for d in discussions:
        if is_undecided(d["content"], d["note"]):
            undecided.append({"type": "discussion", **d})
    for s in issues:
        if is_undecided(s["content"], s["status"], s["plan_date"]):
            undecided.append({"type": "issue", **s})

    result = {
        "source": str(path),
        "sheet": ws.title,
        "overview": read_overview(ws, sec[SEC_DISCUSSION]),
        "discussions": discussions,
        "todos": todos,
        "issues": issues,
        "carryover": {
            "open_todos": open_todos,
            "open_issues": open_issues,
            "undecided": undecided,
        },
        "summary": {
            "discussions": len(discussions),
            "todos": len(todos),
            "issues": len(issues),
            "open_todos": len(open_todos),
            "open_issues": len(open_issues),
            "undecided": len(undecided),
        },
    }

    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
        print(json.dumps({"output": args.output, **result["summary"]},
                         ensure_ascii=False))
    else:
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except AttributeError:
            pass
        print(text)


if __name__ == "__main__":
    main()
